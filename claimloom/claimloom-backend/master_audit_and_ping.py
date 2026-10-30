#!/usr/bin/env python3
"""
ClaimLoom Round 3 — Master File Verifier, Test Asset Generator, Ping & Edge-Case Suite
Run locally:
  cd /Users/pratik/oops-hackathon/claimloom/claimloom-backend
  python3 master_audit_and_ping.py

Run with live Render + Vercel network pings:
  python3 master_audit_and_ping.py --live
"""
from __future__ import annotations

import io
import json
import os
import sys
import time
import zipfile
from pathlib import Path
from typing import Any, Callable, Dict, List, Tuple

_venv_python = Path(__file__).resolve().parent / "venv" / "bin" / "python3"
if _venv_python.is_file() and sys.executable != str(_venv_python.resolve()):
    try:
        import pymupdf
    except ImportError:
        os.execv(str(_venv_python), [str(_venv_python)] + sys.argv)

try:
    import pymupdf as fitz
except ImportError:
    import fitz

import httpx
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from PIL import Image
from starlette.requests import Request

BACKEND_DIR = Path(__file__).resolve().parent
REPO_ROOT = BACKEND_DIR.parent
FRONTEND_DIR = REPO_ROOT / "claimloom-ui"
DATABASE_DIR = REPO_ROOT / "Database"
EXT_DIR = BACKEND_DIR / "chrome-extension"
ASSETS_DIR = BACKEND_DIR / "generated_test_assets"

sys.path.insert(0, str(BACKEND_DIR))

from dotenv import load_dotenv  # noqa: E402
load_dotenv(BACKEND_DIR / ".env")

import ai_client  # noqa: E402
import cache  # noqa: E402
import collab_router  # noqa: E402
import db  # noqa: E402
import main  # noqa: E402
import ratelimit  # noqa: E402

client = TestClient(main.app)


# ==============================================================================
# 1. GENERATE ALL TEST ASSETS (PDFs, PNGs, DOCX, MD, TXT, JSON)
# ==============================================================================
def generate_test_assets() -> Dict[str, bytes]:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    assets: Dict[str, bytes] = {}

    # 01. 2-Page Brochure PDF (Pricing, Guarantee, Terminology, Low Contrast, 6.5pt Font)
    doc = fitz.open()
    p1 = doc.new_page(width=612, height=792)
    p1.insert_text(
        fitz.Point(50, 80),
        "ClaimLoom Pro Plan costs $15/month with guaranteed 100% results for our clients!",
        fontsize=12,
        fontname="helv",
        color=(0.05, 0.05, 0.05),
    )
    p1.insert_text(
        fitz.Point(50, 130),
        "Low contrast subtle disclaimer text on page one",
        fontsize=11,
        fontname="helv",
        color=(0.88, 0.88, 0.88),
    )
    p1.insert_text(
        fitz.Point(50, 720),
        "Unreadable 6.5pt micro-print terms and conditions footer.",
        fontsize=6.5,
        fontname="helv",
        color=(0.10, 0.10, 0.10),
    )
    p2 = doc.new_page(width=612, height=792)
    p2.insert_text(
        fitz.Point(50, 80),
        "Page 2: Certified battery range is 110 km per charge. Click here to start.",
        fontsize=12,
        fontname="helv",
        color=(0.05, 0.05, 0.05),
    )
    assets["01_violation_brochure_2page.pdf"] = doc.tobytes()
    doc.close()

    # 02. Multi-Font PDF (>3 font families -> triggers Rule 9 Consistent Typography)
    doc_fonts = fitz.open()
    pf = doc_fonts.new_page(width=612, height=792)
    for idx, fname in enumerate(("helv", "tiro", "cour", "symb")):
        pf.insert_text(fitz.Point(50, 80 + idx * 30), f"ClaimLoom sample line in {fname}", fontsize=12, fontname=fname)
    assets["02_multifont_brochure.pdf"] = doc_fonts.tobytes()
    doc_fonts.close()

    # 03. Edge Case: Password-Encrypted PDF
    doc_enc = fitz.open()
    doc_enc.new_page()
    assets["03_edge_encrypted.pdf"] = doc_enc.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="secret", user_pw="secret"
    )
    doc_enc.close()

    # 04. Edge Case: 11-Page PDF (> 10-page cap)
    doc_11 = fitz.open()
    for i in range(11):
        pg = doc_11.new_page()
        pg.insert_text(fitz.Point(50, 70), f"Page {i + 1} of 11")
    assets["04_edge_11pages.pdf"] = doc_11.tobytes()
    doc_11.close()

    # 05. Edge Case: Corrupt Non-PDF Header
    assets["05_edge_corrupt.pdf"] = b"NOT_A_VALID_PDF_BINARY"

    # 06. On-Brand High-Contrast PNG (#1E3A8A + #FFFFFF)
    img_ok = Image.new("RGB", (300, 200), (30, 58, 138))
    for x in range(60, 240):
        for y in range(80, 120):
            img_ok.putpixel((x, y), (255, 255, 255))
    b_ok = io.BytesIO()
    img_ok.save(b_ok, format="PNG")
    assets["06_on_brand_banner.png"] = b_ok.getvalue()

    # 07. Off-Brand Neon Lime PNG (#86FA05 -> Rule 9 CIE76 Delta-E > 25)
    img_bad = Image.new("RGB", (300, 200), (134, 250, 5))
    b_bad = io.BytesIO()
    img_bad.save(b_bad, format="PNG")
    assets["07_off_brand_neon.png"] = b_bad.getvalue()

    # 08. Low-Contrast Graphic PNG (#F0F0F0 on #FFFFFF -> Rule 10 Contrast < 4.5:1)
    img_low = Image.new("RGB", (300, 200), (255, 255, 255))
    for x in range(40, 260):
        for y in range(50, 150):
            img_low.putpixel((x, y), (240, 240, 240))
    b_low = io.BytesIO()
    img_low.save(b_low, format="PNG")
    assets["08_low_contrast_graphic.png"] = b_low.getvalue()

    # 09. Transparent RGBA Edge Case PNG ((0,0,0,0) -> composited to White)
    img_rgba = Image.new("RGBA", (100, 100), (0, 0, 0, 0))
    b_rgba = io.BytesIO()
    img_rgba.save(b_rgba, format="PNG")
    assets["09_transparent_rgba_edge.png"] = b_rgba.getvalue()

    # 10. DOCX Document (Paragraph + Table)
    d_obj = Document()
    d_obj.add_heading("ClaimLoom Q4 Launch Brief", level=1)
    d_obj.add_paragraph("Our clients get guaranteed 100% results on the pro plan!")
    tbl = d_obj.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].text = "Plan Tier"
    tbl.rows[0].cells[1].text = "Pro Plan costs $15/month with a lifetime warranty."
    b_docx = io.BytesIO()
    d_obj.save(b_docx)
    assets["10_launch_draft.docx"] = b_docx.getvalue()

    # 11. Markdown Spec Document
    assets["11_spec_sheet.md"] = (
        b"# ClaimLoom Launch Copy\n"
        b"Buy the claim-loom scooter for a 200 km range and zero risk!! Click here now.\n"
    )

    # 12. Edge Case: Empty Text Document
    assets["12_edge_empty.txt"] = b"   \n\t   "

    # 13. Clean 1-Page Compliant PDF
    doc_clean = fitz.open()
    pc = doc_clean.new_page(width=612, height=792)
    pc.insert_text(
        fitz.Point(50, 80),
        "ClaimLoom helps Creators build compliant enterprise assets. The Pro Plan costs $29/month.",
        fontsize=12,
        fontname="helv",
        color=(0.06, 0.09, 0.16),
    )
    assets["13_clean_compliant.pdf"] = doc_clean.tobytes()
    doc_clean.close()

    # 14. JSON Reference File with All Copy-Paste Text Prompts
    prompts = {
        "T01_multi_violation": "Buy the new claim-loom scooter today! It gives a 200 km range with a 100% guarantee.",
        "T02_clean_compliant": "Discover how ClaimLoom helps Creators streamline compliance checks. The Pro Plan costs $29/month and delivers a certified 110 km range under standard testing conditions.",
        "T03_speed_distance_edge": "Top speed is 150km/h and the shop is 200 km away.",
        "T04_singular_plural_edge": "Each client and all clients love our consumer portal for consumers.",
        "T05_emoji_rtl_utf16_edge": "🚀🔥 ClaimLoom Pro Plan is $15/month for مرحبا clients!",
        "T06_pii_redaction_edge": "Mail rohit@gmail.com or call +91 98765 43210, card 4111 1111 1111 1111 for pro plan.",
        "T07_rules_6_8_11_12_13_14_15_16": "Click the green button on the right for our pro plan! It charges to 100% in 10 minutes, includes a lifetime warranty, is the world's best game-changer!! Click here.",
        "C01_billboard": "Visit https://claimloom.ai or click here right now while driving on the highway today.",
        "C02_linkedin": "AMAZING update!! DM me for details #ai #saas #tech #b2b #growth #sales",
        "C03_instagram": "See https://claimloom.ai/launch",
        "C04_email": "ACT NOW! This is 100% free with FREE MONEY TODAY ONLY for subscribers.",
    }
    assets["14_copy_paste_text_cases.json"] = json.dumps(prompts, indent=2, ensure_ascii=False).encode("utf-8")

    for name, data in assets.items():
        (ASSETS_DIR / name).write_bytes(data)

    return assets


# ==============================================================================
# 2. CHECK ALL REPOSITORY FILES (BACKEND, FRONTEND, EXTENSION, DATABASE)
# ==============================================================================
PROJECT_FILES: List[Tuple[Path, str]] = [
    (BACKEND_DIR, "main.py"),
    (BACKEND_DIR, "ai_client.py"),
    (BACKEND_DIR, "benchmark.py"),
    (BACKEND_DIR, "cache.py"),
    (BACKEND_DIR, "collab_router.py"),
    (BACKEND_DIR, "db.py"),
    (BACKEND_DIR, "engine.py"),
    (BACKEND_DIR, "kb_router.py"),
    (BACKEND_DIR, "pdf_engine.py"),
    (BACKEND_DIR, "pipeline.py"),
    (BACKEND_DIR, "rag_verifier.py"),
    (BACKEND_DIR, "ratelimit.py"),
    (BACKEND_DIR, "rule_version.py"),
    (BACKEND_DIR, "rules_router.py"),
    (BACKEND_DIR, "schemas.py"),
    (BACKEND_DIR, "seed_db.py"),
    (BACKEND_DIR, "seed_demo_scans.py"),
    (BACKEND_DIR, "seed_kb.py"),
    (BACKEND_DIR, "session.py"),
    (BACKEND_DIR, "signing.py"),
    (BACKEND_DIR, "vision_router.py"),
    (BACKEND_DIR, "webhook_router.py"),
    (BACKEND_DIR, "requirements.txt"),
    (BACKEND_DIR, "migrations/001_round3.sql"),
    (BACKEND_DIR, "tests/test_backend.py"),
    (BACKEND_DIR, "tests/test_round3.py"),
    (BACKEND_DIR, "verify_backend.py"),
    (EXT_DIR, "manifest.json"),
    (EXT_DIR, "popup.html"),
    (EXT_DIR, "popup.css"),
    (EXT_DIR, "popup.js"),
    (FRONTEND_DIR, "package.json"),
    (FRONTEND_DIR, "postcss.config.mjs"),
    (FRONTEND_DIR, "next.config.ts"),
    (FRONTEND_DIR, "tsconfig.json"),
    (FRONTEND_DIR, "app/globals.css"),
    (FRONTEND_DIR, "app/layout.tsx"),
    (FRONTEND_DIR, "app/page.tsx"),
    (FRONTEND_DIR, "components/CanvasTab.tsx"),
    (FRONTEND_DIR, "components/CertificateModal.tsx"),
    (FRONTEND_DIR, "components/Common.tsx"),
    (FRONTEND_DIR, "components/EditorCard.tsx"),
    (FRONTEND_DIR, "components/ExtensionModal.tsx"),
    (FRONTEND_DIR, "components/RadarTab.tsx"),
    (FRONTEND_DIR, "components/ResultsPanel.tsx"),
    (FRONTEND_DIR, "components/SignoffCard.tsx"),
    (FRONTEND_DIR, "components/VaultTab.tsx"),
    (FRONTEND_DIR, "lib/api.ts"),
    (FRONTEND_DIR, "lib/client.ts"),
    (FRONTEND_DIR, "lib/mockData.ts"),
    (FRONTEND_DIR, "lib/presets.ts"),
    (FRONTEND_DIR, "lib/schema.ts"),
    (FRONTEND_DIR, "lib/ui.ts"),
    (FRONTEND_DIR, "lib/useClaimLoomWorkspace.ts"),
    (DATABASE_DIR, "test_setup.py"),
    (DATABASE_DIR, "test_embedding.py"),
]


def check_all_files() -> Tuple[int, int]:
    print("\n" + "=" * 100)
    print(" STEP 1: REPOSITORY FILE VERIFICATION & SYNTAX CHECK")
    print("=" * 100)
    ok_count = 0
    for folder, rel in PROJECT_FILES:
        p = folder / rel
        label = f"{folder.name}/{rel}"
        if not p.is_file() or p.stat().st_size == 0:
            print(f"  [MISSING] {label}")
            continue
        try:
            if p.suffix == ".py":
                compile(p.read_text(encoding="utf-8"), str(p), "exec")
            elif p.suffix == ".json":
                json.loads(p.read_text(encoding="utf-8"))
            ok_count += 1
            print(f"  [FILE OK] {label:<56} | {p.stat().st_size:>8,} bytes")
        except Exception as exc:
            print(f"  [SYNTAX!] {label:<56} | {exc}")
    print(f"\n  -> Verified {ok_count}/{len(PROJECT_FILES)} repository files.")
    return ok_count, len(PROJECT_FILES)


# ==============================================================================
# 3. ENDPOINT PING & 42-CASE NORMAL + EDGE-CASE VERIFICATION SUITE
# ==============================================================================
def run_all_tests(assets: Dict[str, bytes], live_ping: bool = False) -> None:
    print("\n" + "=" * 100)
    print(" STEP 2: ENDPOINT PINGS & 42 NORMAL + EDGE-CASE ASSERTIONS")
    print("=" * 100)

    ratelimit.reset_rate_limits()
    orig_mode = ai_client.AI_MODE
    os.environ["AI_MODE"] = "rules_only"
    ai_client.AI_MODE = "rules_only"

    results: List[Tuple[str, bool, float, str]] = []

    def check(code: str, name: str, fn: Callable[[], str]) -> None:
        t0 = time.perf_counter()
        try:
            info = fn()
            ms = (time.perf_counter() - t0) * 1000.0
            print(f"  [PASS] {code:<4} | {name:<48} | {ms:6.1f} ms | {info}")
            results.append((f"{code} {name}", True, ms, info))
        except Exception as exc:
            ms = (time.perf_counter() - t0) * 1000.0
            print(f"  [FAIL] {code:<4} | {name:<48} | {ms:6.1f} ms | {type(exc).__name__}: {exc}")
            results.append((f"{code} {name}", False, ms, str(exc)))

    # --- PINGS & DIAGNOSTICS ---
    check("P01", "GET /ping (Keep-Alive Ping)", lambda: (
        (r := client.get("/ping")),
        assert_true(r.status_code == 200 and r.json().get("status") == "awake", r.text),
        f"HTTP 200 ({r.json().get('service')})"
    )[-1])

    check("P02", "GET /health (Database & Vault Ping)", lambda: (
        (r := client.get("/health")),
        assert_true(r.status_code == 200 and "rules_source" in r.json(), r.text),
        f"HTTP 200 (source={r.json()['rules_source']})"
    )[-1])

    check("P03", "GET /api/public-key (Ed25519 Key Ping)", lambda: (
        (r := client.get("/api/public-key")),
        assert_true(r.status_code == 200 and r.json().get("algorithm") == "Ed25519", r.text),
        f"HTTP 200 (key={r.json()['public_key'][:16]}...)"
    )[-1])

    check("P04", "GET /api/rules & /api/kb/sources Ping", lambda: (
        (r1 := client.get("/api/rules?refresh=true")),
        (r2 := client.get("/api/kb/sources?workspace_id=default")),
        assert_true(r1.status_code == 200 and len(r1.json()["rules"]) >= 16, r1.text),
        assert_true(r2.status_code == 200 and len(r2.json()["chunks"]) >= 13, r2.text),
        f"rules={len(r1.json()['rules'])}, chunks={len(r2.json()['chunks'])}"
    )[-1])

    # --- SESSIONS & RBAC ---
    tokens: Dict[str, str] = {}

    def _setup_sessions() -> str:
        for role in ("Creator", "Reviewer", "Admin"):
            r = client.post("/api/session", json={"name": f"Test {role}", "role": role, "workspace_id": "default"})
            assert_true(r.status_code == 200, r.text)
            tokens[role] = r.json()["token"]
        bad_role = client.post("/api/session", json={"name": "Bad", "role": "SuperUser", "workspace_id": "default"})
        assert_true(bad_role.status_code == 400, "Expected 400 on invalid role")
        bad_token = client.get("/api/session/me", headers={"Authorization": "Bearer tampered.token"})
        assert_true(bad_token.status_code == 401, "Expected 401 on invalid token")
        ws_mismatch = client.get(
            "/api/session/me",
            headers={"Authorization": f"Bearer {tokens['Creator']}", "X-Workspace-Id": "custom_brand"},
        )
        assert_true(ws_mismatch.status_code == 403, "Expected 403 on workspace header mismatch")
        return "Tokens issued + invalid role(400), bad token(401), ws mismatch(403) verified"

    check("S01", "Session Issuance + Auth Edge Cases (400/401/403)", _setup_sessions)

    # Ensure clean baseline demo state
    client.post("/api/admin/reset-demo")
    ratelimit.reset_rate_limits()

    # --- TEXT ENGINE NORMAL & EDGE CASES ---
    def _t01() -> str:
        txt = "Buy the new claim-loom scooter today! It gives a 200 km range with a 100% guarantee."
        r = client.post("/api/scan-text", json={"text": txt, "channel": "general"})
        d = r.json()
        rids = {i["rule_id"] for i in d["issues"]}
        assert_true(rids == {1, 2, 7} and d["readiness_score"] == 48, f"Got {rids}, score={d['readiness_score']}")
        for i in d["issues"]:
            assert_true(txt[i["start"]:i["end"]] == i["flagged_text"], f"Offset mismatch on {i}")
        return f"score=48, rules={sorted(rids)}, exact offsets verified"

    def _t02() -> str:
        txt = (
            "Discover how ClaimLoom helps Creators streamline compliance checks. "
            "The Pro Plan costs $29/month and delivers a certified 110 km range under standard testing conditions."
        )
        d = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
        assert_true(d["readiness_score"] == 100 and d["issues"] == [], f"Unexpected flags: {d['issues']}")
        assert_true(len(d["verified_claims"]) >= 2, "Expected >=2 verified RAG claims")
        return f"score=100, flags=0, verified_claims={len(d['verified_claims'])}"

    def _t03() -> str:
        txt = "Top speed is 150km/h and the shop is 200 km away."
        d = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
        assert_true(d["issues"] == [] and d["readiness_score"] == 100, f"False positive: {d['issues']}")
        return "0 false positives on speed (150km/h) or distance (200 km away)"

    def _t04() -> str:
        txt = "Each client and all clients love our consumer portal for consumers."
        d = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
        fixes = [i["suggested_fix"] for i in d["issues"] if i["rule_id"] == 3]
        assert_true(fixes == ["Creator", "Creators", "Creator", "Creators"], f"Got {fixes}")
        return f"singular/plural fixes={fixes}"

    def _t05() -> str:
        txt = "🚀🔥 ClaimLoom Pro Plan is $15/month for مرحبا clients!"
        d = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
        p_iss = next(i for i in d["issues"] if i["rule_id"] == 5)
        c_iss = next(i for i in d["issues"] if i["rule_id"] == 3)
        assert_true((p_iss["start"], p_iss["end"]) == (27, 36), f"Got {p_iss}")
        assert_true((c_iss["start"], c_iss["end"]) == (47, 54), f"Got {c_iss}")
        return f"UTF-16 offsets: $15/month=[27,36], clients=[47,54]"

    def _t06() -> str:
        txt = "Mail rohit@gmail.com or call +91 98765 43210, card 4111 1111 1111 1111 for pro plan."
        d = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
        assert_true(d["pii_redacted_count"] == 3, f"Expected 3 PII redactions, got {d['pii_redacted_count']}")
        assert_true("rohit@gmail.com" not in d["redacted_text"] and len(d["redacted_text"]) == len(txt), "PII leak")
        r11 = next(i for i in d["issues"] if i["rule_id"] == 11)
        assert_true(txt[r11["start"]:r11["end"]] == "pro plan", "Offset shifted after PII redaction")
        return "3 PII items masked at equal length; 'pro plan' offset [75,83] intact"

    def _t07() -> str:
        txt = (
            "Click the green button on the right for our pro plan! It charges to 100% in 10 minutes, "
            "includes a lifetime warranty, is the world's best game-changer!! Click here."
        )
        d = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
        rids = {i["rule_id"] for i in d["issues"]}
        exp = {6, 8, 11, 12, 13, 14, 15, 16}
        assert_true(exp.issubset(rids), f"Missing rules: {exp - rids}")
        return f"rules {sorted(exp)} all flagged, score={d['readiness_score']}"

    check("T01", "Multi-Rule Violation Copy (Rules 1, 2, 7)", _t01)
    check("T02", "Compliant Copy + No-Op Guard + Verified RAG", _t02)
    check("T03", "Edge: Speed (150km/h) & Distance (200 km away)", _t03)
    check("T04", "Edge: Singular/Plural Dynamic Fix (Rule 3)", _t04)
    check("T05", "Edge: Astral Emoji + RTL Arabic UTF-16 Offsets", _t05)
    check("T06", "Edge: Same-Length PII Redaction (Email/Tel/Card)", _t06)
    check("T07", "Sweep Rules #6, #8, #11, #12, #13, #14, #15, #16", _t07)
    check("T08", "Edge: Empty/Whitespace Text Rejection (400)", lambda: (
        assert_true(client.post("/api/scan-text", json={"text": "   \n\t "}).status_code == 400, "Expected 400"),
        "HTTP 400 Bad Request"
    )[-1])
    check("T09", "Edge: Oversized Text >8000 Chars Rejection (422)", lambda: (
        assert_true(client.post("/api/scan-text", json={"text": "A" * 8001}).status_code == 422, "Expected 422"),
        "HTTP 422 Unprocessable Entity"
    )[-1])

    # --- CHANNEL PLAYBOOKS ---
    def _c01_channels() -> str:
        bb = client.post("/api/scan-text", json={
            "text": "Visit https://claimloom.ai or click here right now while driving on the highway today.",
            "channel": "billboard"
        }).json()
        assert_true({1003, 1006}.issubset({i["rule_id"] for i in bb["issues"]}), "Billboard failed")

        li = client.post("/api/scan-text", json={
            "text": "AMAZING update!! DM me for details #ai #saas #tech #b2b #growth #sales",
            "channel": "linkedin"
        }).json()
        assert_true({1001, 1007, 1008, 1009}.issubset({i["rule_id"] for i in li["issues"]}), "LinkedIn failed")

        ig = client.post("/api/scan-text", json={
            "text": "See https://claimloom.ai/launch " + ("word " * 152),
            "channel": "instagram"
        }).json()
        assert_true({1002, 1010}.issubset({i["rule_id"] for i in ig["issues"]}), "Instagram failed")

        em = client.post("/api/scan-text", json={
            "text": "ACT NOW! This is 100% free with FREE MONEY TODAY ONLY for subscribers.",
            "channel": "email"
        }).json()
        assert_true({1004, 1005}.issubset({i["rule_id"] for i in em["issues"]}), "Email failed")
        return "All 10 channel-specific rules (1001..1010) verified"

    check("C01", "Channel Playbooks (Billboard/LinkedIn/IG/Email)", _c01_channels)

    # --- AI TIER FAILOVER, JSON REPAIR & HALLUCINATION FILTER ---
    def _ai_failover_and_guards() -> str:
        ratelimit.reset_rate_limits()
        ai_client.reset_cooldowns()
        os.environ["GROQ_API_KEY"] = "gsk_simulated"
        os.environ["GEMINI_API_KEY"] = "gem_simulated"
        ai_client.AI_MODE = "auto"
        orig_post = ai_client._post
        calls: List[str] = []
        try:
            def mock_post(url: str, headers: dict, payload: dict, timeout: float) -> Tuple[int, str]:
                calls.append(url)
                if "groq" in url:
                    return 404, '{"error":"model retired"}'
                # Return 1 valid Rule 4 candidate + 1 hallucinated rule 99 + 1 no-op self-fix
                cands = [
                    {"rule_id": 4, "flagged_text": "crush and humiliate every rival", "explanation": "Aggressive tone.", "suggested_fix": "lead our category", "confidence": 0.91},
                    {"rule_id": 99, "flagged_text": "ClaimLoom", "explanation": "Hallucinated rule", "suggested_fix": "Other"},
                    {"rule_id": 1, "flagged_text": "ClaimLoom", "explanation": "No-op self fix", "suggested_fix": "ClaimLoom"},
                ]
                return 200, json.dumps({"choices": [{"message": {"content": json.dumps({"candidates": cands})}}]})

            ai_client._post = mock_post
            txt = "ClaimLoom will crush and humiliate every rival in the market."
            r1 = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
            assert_true(r1["ai_status"] == "ok" and r1["cached"] is False, f"Got {r1}")
            assert_true([i["rule_id"] for i in r1["issues"]] == [4], f"Got issues: {r1['issues']}")

            # Second identical scan must hit SHA-256 cache (cached=True)
            r2 = client.post("/api/scan-text", json={"text": txt, "channel": "general"}).json()
            assert_true(r2["cached"] is True, "Expected SHA-256 cache hit on second identical scan")
            return f"Groq(404)->Gemini(200) failover + hallucination filter + SHA-256 cache(True)"
        finally:
            ai_client._post = orig_post
            ai_client.AI_MODE = "rules_only"
            os.environ.pop("GROQ_API_KEY", None)
            os.environ.pop("GEMINI_API_KEY", None)

    check("A01", "AI Provider Failover + Hallucination Guard + Cache", _ai_failover_and_guards)

    # --- IMAGE SCANS & EDGE CASES ---
    def _image_suite() -> str:
        ratelimit.reset_rate_limits()
        # 06. On-brand PNG
        r_ok = client.post("/api/scan-image", files={"file": ("06.png", assets["06_on_brand_banner.png"], "image/png")}).json()
        assert_true(r_ok["issues"] == [], f"Unexpected issues on brand PNG: {r_ok['issues']}")
        # Verify Canvas PNG & Redline PDF work on uploaded images
        assert_true(client.get(f"/api/scans/{r_ok['scan_id']}/page/1.png").status_code == 200, "Image page/1.png failed")
        assert_true(client.get(f"/api/scans/{r_ok['scan_id']}/redline-pdf").status_code == 200, "Image redline-pdf failed")

        # 07. Off-brand neon lime PNG -> Rule 9
        r_neon = client.post("/api/scan-image", files={"file": ("07.png", assets["07_off_brand_neon.png"], "image/png")}).json()
        assert_true(9 in [i["rule_id"] for i in r_neon["issues"]], "Expected Rule 9 on neon lime image")

        # 08. Low-contrast graphic PNG -> Rule 10
        r_low = client.post("/api/scan-image", files={"file": ("08.png", assets["08_low_contrast_graphic.png"], "image/png")}).json()
        assert_true(10 in [i["rule_id"] for i in r_low["issues"]], "Expected Rule 10 on low-contrast image")

        # 09. Transparent RGBA PNG -> composited onto white, 0 black false positives
        r_rgba = client.post("/api/scan-image", files={"file": ("09.png", assets["09_transparent_rgba_edge.png"], "image/png")}).json()
        assert_true(not any("#000000" in i["flagged_text"] for i in r_rgba["issues"]), "RGBA transparency bug")

        # Corrupt / non-image uploads -> HTTP 400
        assert_true(client.post("/api/scan-image", files={"file": ("bad.png", b"not_an_image", "image/png")}).status_code == 400, "Expected 400")
        assert_true(client.post("/api/scan-image", files={"file": ("bad.txt", b"hello", "text/plain")}).status_code == 400, "Expected 400")
        return "On-brand(0), Off-brand(Rule 9), Low-contrast(Rule 10), RGBA(OK), Corrupt(400)"

    check("I01", "Image Suite (CIE76, Contrast, RGBA, Canvas PNG)", _image_suite)

    # --- PDF & DOCUMENT SCANS + EDGE CASES ---
    def _pdf_suite() -> str:
        ratelimit.reset_rate_limits()
        # 01. 2-page violation brochure
        r_pdf = client.post("/api/scan-pdf", files={"file": ("01.pdf", assets["01_violation_brochure_2page.pdf"], "application/pdf")})
        assert_true(r_pdf.status_code == 200, r_pdf.text)
        d = r_pdf.json()
        assert_true(d["page_count"] == 2 and any(i["detected_by"] == "pdf_layout" for i in d["issues"]), "Missing pdf_layout")
        for iss in d["issues"]:
            assert_true(iss["page"] in (1, 2) and len(iss["bbox"]) == 4, f"Bad bbox: {iss}")

        # Page PNG & Redlined 3-Page PDF
        p1_png = client.get(f"/api/scans/{d['scan_id']}/page/1.png")
        p99_png = client.get(f"/api/scans/{d['scan_id']}/page/99.png")
        assert_true(p1_png.status_code == 200 and p1_png.content.startswith(b"\x89PNG"), "Page 1 PNG failed")
        assert_true(p99_png.status_code == 404, "Page 99 out-of-range should return 404")

        red = client.get(f"/api/scans/{d['scan_id']}/redline-pdf")
        assert_true(red.status_code == 200 and red.content.startswith(b"%PDF"), "Redline PDF failed")
        with fitz.open(stream=red.content, filetype="pdf") as rdoc:
            assert_true(len(rdoc) == 3 and "CRYPTOGRAPHIC COMPLIANCE & AUDIT SEAL" in rdoc[-1].get_text("text"), "Cert page missing")

        # 02. Multi-font PDF (>3 fonts -> Rule 9)
        r_mf = client.post("/api/scan-pdf", files={"file": ("02.pdf", assets["02_multifont_brochure.pdf"], "application/pdf")}).json()
        assert_true(any(i["rule_id"] == 9 and i["detected_by"] == "pdf_layout" for i in r_mf["issues"]), "Expected >3 font family flag")

        # 03, 04, 05. Edge Case PDFs (Encrypted, 11-Page, Corrupt)
        assert_true(client.post("/api/scan-pdf", files={"file": ("03.pdf", assets["03_edge_encrypted.pdf"], "application/pdf")}).status_code == 400, "Encrypted PDF should 400")
        assert_true(client.post("/api/scan-pdf", files={"file": ("04.pdf", assets["04_edge_11pages.pdf"], "application/pdf")}).status_code == 400, "11-page PDF should 400")
        assert_true(client.post("/api/scan-pdf", files={"file": ("05.pdf", assets["05_edge_corrupt.pdf"], "application/pdf")}).status_code == 400, "Corrupt PDF should 400")
        return "2-page BBox+Layout OK, Multi-font(Rule 9) OK, Encrypted/11p/Corrupt(400) OK"

    def _doc_suite() -> str:
        ratelimit.reset_rate_limits()
        r_docx = client.post("/api/scan-doc", files={"file": ("10.docx", assets["10_launch_draft.docx"], "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert_true(r_docx.status_code == 200, r_docx.text)
        assert_true({3, 5, 7, 11, 12}.issubset({i["rule_id"] for i in r_docx.json()["issues"]}), "DOCX table/paragraph scan incomplete")

        r_md = client.post("/api/scan-doc", files={"file": ("11.md", assets["11_spec_sheet.md"], "text/markdown")})
        assert_true(r_md.status_code == 200 and {1, 2, 7, 14, 16}.issubset({i["rule_id"] for i in r_md.json()["issues"]}), "MD scan incomplete")

        # Edge cases: empty .txt (400), unsupported .exe (400)
        assert_true(client.post("/api/scan-doc", files={"file": ("12.txt", assets["12_edge_empty.txt"], "text/plain")}).status_code == 400, "Empty doc should 400")
        assert_true(client.post("/api/scan-doc", files={"file": ("bad.exe", b"binary", "application/octet-stream")}).status_code == 400, "Unsupported ext should 400")
        return "DOCX(paragraphs+tables) OK, Markdown OK, Empty/Unsupported(400) OK"

    check("F01", "PDF Suite (BBox, Font/Contrast, Redline, Edges)", _pdf_suite)
    check("D01", "Document Suite (DOCX, Markdown, Empty/Bad Ext)", _doc_suite)

    # --- TRUTH VAULT PRICE SYNC ($29 -> $39) & RE-AUDIT CASCADE ---
    def _vault_sync() -> str:
        ratelimit.reset_rate_limits()
        up = client.put("/api/kb/chunks/3", json={"chunk_text": "ClaimLoom Pro Plan subscription price is $39/month or $390/year."})
        assert_true(up.status_code == 200 and up.json()["affected_count"] >= 2, f"Got {up.text}")
        r39 = client.post("/api/scan-text", json={"text": "The Pro Plan costs $39/month."}).json()
        r29 = client.post("/api/scan-text", json={"text": "The Pro Plan costs $29/month."}).json()
        assert_true(not any(i["rule_id"] == 5 for i in r39["issues"]), "$39/month should pass after sync")
        assert_true(any(i["rule_id"] == 5 for i in r29["issues"]), "$29/month should flag after sync")
        client.post("/api/admin/reset-demo")
        return f"affected={up.json()['affected_count']}, $39=Pass, $29=Flagged, Reset=Restored"

    check("K01", "Truth Vault Price Sync ($29->$39) & Rule #5 Sync", _vault_sync)

    # --- COLLABORATION, ED25519 SIGNING, TAMPER DETECTION & EXTENSION ZIP ---
    def _collab_and_signing() -> str:
        ratelimit.reset_rate_limits()
        scan = client.post("/api/scan-text", json={"text": "ClaimLoom Pro Plan is $15/month with 100% guarantee."}).json()
        sid = scan["scan_id"]

        # Comment creation & activity feed
        c_res = client.post(f"/api/scans/{sid}/comments", json={"comment_text": "Update price to $29/month."}, headers={"Authorization": f"Bearer {tokens['Reviewer']}"})
        assert_true(c_res.status_code == 200, c_res.text)
        act = client.get(f"/api/scans/{sid}/activity", headers={"Authorization": f"Bearer {tokens['Reviewer']}"}).json()
        assert_true(any(a["type"] == "comment" for a in act["activity"]), "Comment missing from activity")

        # High-severity blocker gate -> HTTP 409
        b_res = client.post(f"/api/scans/{sid}/approve", json={}, headers={"Authorization": f"Bearer {tokens['Reviewer']}"})
        assert_true(b_res.status_code == 409, f"Expected 409 blocker gate, got {b_res.status_code}")

        # Heal copy in Version 2 -> Approve -> Verify Ed25519 chain
        v2 = client.post(f"/api/scans/{sid}/versions", json={"text_snapshot": "ClaimLoom Pro Plan is $29/month with proven results."}, headers={"Authorization": f"Bearer {tokens['Creator']}"})
        assert_true(v2.status_code == 200 and v2.json()["version_no"] == 2, v2.text)

        app_res = client.post(f"/api/scans/{sid}/approve", json={"version_no": 2}, headers={"Authorization": f"Bearer {tokens['Reviewer']}"})
        assert_true(app_res.status_code == 200 and app_res.json()["approval"]["signature"], app_res.text)

        ver1 = client.get(f"/api/verify/{sid}").json()
        assert_true(ver1["valid"] is True and ver1["tampered"] is False and ver1["chain_valid"] is True, f"Verify failed: {ver1}")

        # Post-approval tamper -> Version 3
        client.post(f"/api/scans/{sid}/versions", json={"text_snapshot": "Sneaky post-approval edit $1/month!"}, headers={"Authorization": f"Bearer {tokens['Creator']}"})
        ver2 = client.get(f"/api/verify/{sid}").json()
        assert_true(ver2["valid"] is False and ver2["tampered"] is True, "Tamper detection failed")

        # Chrome Extension ZIP bundle
        z_res = client.get("/api/extension/download")
        assert_true(z_res.status_code == 200 and z_res.content.startswith(b"PK"), "Extension ZIP failed")
        with zipfile.ZipFile(io.BytesIO(z_res.content)) as zf:
            assert_true({"manifest.json", "popup.html", "popup.css", "popup.js"}.issubset(set(zf.namelist())), "ZIP missing files")
        return "Comments OK, 409 Gate OK, v2 Ed25519 Sign OK, v3 Tamper Caught, Extension ZIP OK"

    check("G01", "Collab Thread, Ed25519 Sign, Tamper & Ext ZIP", _collab_and_signing)

    # --- SELF-HEALING RULES, REDOS AST GUARD, WEBHOOK & RATE LIMITER ---
    def _security_and_radar() -> str:
        admin_h = {"Authorization": f"Bearer {tokens['Admin']}"}
        creator_h = {"Authorization": f"Bearer {tokens['Creator']}"}

        # Flag action override requires reason (422 without reason, 200 with reason)
        assert_true(client.post("/api/flag-action", json={"scan_id": "s1", "rule_id": 3, "action": "Overridden", "reason": ""}).status_code == 422, "Override without reason must 422")
        assert_true(client.post("/api/flag-action", json={"scan_id": "s1", "rule_id": 3, "action": "Overridden", "reason": "B2B contract quote"}).status_code == 200, "Override with reason failed")

        # ReDoS guard (422), Self-Fix match guard (422), Broad regex guard (422), RBAC (403)
        assert_true(client.post("/api/rules/3/backtest", json={"proposed_regex": r"(a+)+$"}, headers=admin_h).status_code == 422, "ReDoS must 422")
        assert_true(client.post("/api/rules/1/backtest", json={"proposed_regex": r"(?i)\bClaimLoom\b"}, headers=admin_h).status_code == 422, "Self-fix must 422")
        assert_true(client.post("/api/rules", json={"violation_type": "Tone", "severity": "Low", "rule_cited": "Broad", "rule_text": "Broad test", "regex_pattern": ".*"}, headers=admin_h).status_code == 422, "Broad .* must 422")
        assert_true(client.patch("/api/rules/3", json={"rule_text": "Creator edit"}, headers=creator_h).status_code == 403, "Creator PATCH must 403")

        # Webhook auth check
        os.environ["WEBHOOK_SECRET"] = "top_secret_hook"
        try:
            ratelimit.reset_rate_limits()
            assert_true(client.post("/api/webhooks/scan", json={"text": "Hello world."}, headers={"X-Webhook-Secret": "wrong"}).status_code == 401, "Webhook wrong secret must 401")
            assert_true(client.post("/api/webhooks/scan", json={"text": "ClaimLoom Pro Plan is $29/month."}, headers={"X-Webhook-Secret": "top_secret_hook"}).status_code == 200, "Webhook valid secret must 200")
        finally:
            os.environ.pop("WEBHOOK_SECRET", None)

        # Proxy X-Forwarded-For Rate Limiter
        ratelimit.reset_rate_limits()
        req = Request({
            "type": "http", "http_version": "1.1", "method": "POST", "scheme": "http",
            "path": "/api/scan-text", "raw_path": b"/api/scan-text", "query_string": b"",
            "headers": [(b"x-forwarded-for", b"10.0.0.1, 198.51.100.42")],
            "client": ("127.0.0.1", 9000), "server": ("testserver", 80),
        })
        for _ in range(int(ratelimit.BUCKET_CAPACITY)):
            assert_true(ratelimit.rate_limit_scan(req) == "198.51.100.42", "XFF IP mismatch")
        try:
            ratelimit.rate_limit_scan(req)
            raise AssertionError("Expected HTTP 429")
        except HTTPException as exc:
            assert_true(exc.status_code == 429 and "Retry-After" in exc.headers, "Expected 429 + Retry-After")
        finally:
            ratelimit.reset_rate_limits()

        return "Override(422/200), ReDoS(422), RBAC(403), Webhook(401/200), RateLimit(429) OK"

    check("R01", "ReDoS Guard, RBAC, Webhook Auth & Rate Limiter", _security_and_radar)

    ai_client.AI_MODE = orig_mode

    # --- OPTIONAL LIVE NETWORK PINGS ---
    if live_ping:
        print("\n" + "=" * 100)
        print(" STEP 3: LIVE DEPLOYMENT NETWORK PINGS (RENDER & VERCEL)")
        print("=" * 100)
        endpoints = [
            ("Render GET /ping", "https://oops-hackathon.onrender.com/ping"),
            ("Render GET /health", "https://oops-hackathon.onrender.com/health"),
            ("Render GET /api/rules", "https://oops-hackathon.onrender.com/api/rules"),
            ("Render GET /api/kb/sources", "https://oops-hackathon.onrender.com/api/kb/sources"),
            ("Render GET /api/extension/download", "https://oops-hackathon.onrender.com/api/extension/download"),
            ("Vercel Frontend UI", "https://oops-hackathon.vercel.app"),
        ]
        for label, url in endpoints:
            t0 = time.perf_counter()
            try:
                r = httpx.get(url, timeout=25.0, follow_redirects=True)
                ms = (time.perf_counter() - t0) * 1000.0
                print(f"  [LIVE OK] {label:<36} | HTTP {r.status_code} | {ms:7.1f} ms | {len(r.content):>8,} bytes")
            except Exception as exc:
                ms = (time.perf_counter() - t0) * 1000.0
                print(f"  [FAILED!] {label:<36} | ERROR    | {ms:7.1f} ms | {exc}")

    print("\n" + "=" * 100)
    passed = sum(1 for _, ok, _, _ in results if ok)
    print(f" AUDIT RESULT: {passed}/{len(results)} TEST GROUPS PASSED | GENERATED 14 TEST FILES IN: {ASSETS_DIR}")
    print("=" * 100 + "\n")
    if passed < len(results):
        raise SystemExit(1)


def assert_true(cond: Any, msg: str = "Assertion failed") -> None:
    if not cond:
        raise AssertionError(msg)


if __name__ == "__main__":
    generated_assets = generate_test_assets()
    check_all_files()
    run_all_tests(generated_assets, live_ping=("--live" in sys.argv))
